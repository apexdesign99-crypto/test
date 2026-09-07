"""テスト用の共通部品(API を呼ばないダミークライアント)。"""

from __future__ import annotations

from types import SimpleNamespace
from typing import Any

import pytest

from ai_employee.profile import build_profile
from ai_employee.workspace import Workspace


def text_delta(text: str) -> SimpleNamespace:
    return SimpleNamespace(
        type="content_block_delta", delta=SimpleNamespace(type="text_delta", text=text)
    )


def thinking_delta(text: str) -> SimpleNamespace:
    return SimpleNamespace(
        type="content_block_delta",
        delta=SimpleNamespace(type="thinking_delta", thinking=text),
    )


def message(
    content: list[dict[str, Any]],
    stop_reason: str = "end_turn",
    stop_details: Any = None,
) -> SimpleNamespace:
    return SimpleNamespace(
        content=content, stop_reason=stop_reason, stop_details=stop_details
    )


def say(text: str, stop_reason: str = "end_turn") -> SimpleNamespace:
    return message([{"type": "text", "text": text}], stop_reason)


def call_tool(name: str, arguments: dict[str, Any], tool_id: str = "tu_1"):
    return message(
        [{"type": "tool_use", "id": tool_id, "name": name, "input": arguments}],
        stop_reason="tool_use",
    )


class _FakeStream:
    def __init__(self, final: SimpleNamespace) -> None:
        self._final = final
        self.closed = False

    def __enter__(self) -> "_FakeStream":
        return self

    def __exit__(self, *exc: object) -> bool:
        self.closed = True
        return False

    def __iter__(self):
        for block in self._final.content:
            if block.get("type") == "text":
                yield text_delta(block["text"])
            elif block.get("type") == "thinking":
                yield thinking_delta(block.get("thinking", ""))

    def get_final_message(self) -> SimpleNamespace:
        return self._final


class FakeMessages:
    """台本どおりに応答を返す `client.beta.messages` の代役。"""

    def __init__(self, script: list[SimpleNamespace]) -> None:
        self.script = list(script)
        self.calls: list[dict[str, Any]] = []
        self.streams: list[_FakeStream] = []

    def stream(self, **kwargs: Any) -> _FakeStream:
        self.calls.append(kwargs)
        if not self.script:
            raise AssertionError("台本より多く API が呼ばれました")
        stream = _FakeStream(self.script.pop(0))
        self.streams.append(stream)
        return stream


class FakeClient:
    def __init__(self, script: list[SimpleNamespace]) -> None:
        self.messages = FakeMessages(script)
        self.beta = SimpleNamespace(messages=self.messages)


@pytest.fixture
def workspace(tmp_path) -> Workspace:
    ws = Workspace("tester", tmp_path)
    ws.save_profile(build_profile("tester", "テスト社員", "assistant"))
    return ws


@pytest.fixture
def profile(workspace):
    return workspace.load_profile()


@pytest.fixture
def テンプレート(tmp_path):
    """事務所のプラン台帳を複製し、事例を空にして渡す。

    数式・入力規則は本物でないと意味がないので簡易版では代用しない。
    一方、出荷ファイルに入っている事例はテストの前提にしない——
    拾い出しをするたびにテストが落ちることになる。
    """
    import shutil

    from ai_employee.plans import DEFAULT_DB

    pytest.importorskip("openpyxl")
    if not DEFAULT_DB.is_file():
        pytest.skip("テンプレートが見つかりません")
    path = tmp_path / "db.xlsx"
    shutil.copy(DEFAULT_DB, path)

    from openpyxl import load_workbook

    workbook = load_workbook(path)  # data_only=False——数式は残す
    for sheet in workbook.worksheets:
        if not sheet.title.startswith(("①", "②", "③", "④", "⑤", "⑥")):
            continue
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if not (isinstance(cell.value, str) and cell.value.startswith("=")):
                    cell.value = None
    workbook.save(path)
    return path
